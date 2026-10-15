# validar.py — Etapa 1: validar el cruce de datos (sin web todavía).
# Uso cuando haya Python: pip install openpyxl ; python validar.py
# Lee INSUMOS/*.xlsx (solo lectura) y muestra resumen + tabla + inconsistencias.

from openpyxl import load_workbook  # librería para leer Excel .xlsx
from collections import defaultdict


# LADRILLO función: agrupa una tarea con nombre para reutilizarla.
def calcular_promedio(numeros):
    # LADRILLO variables y tipos: lista de números (int/float) -> float.
    if not numeros:  # LADRILLO condicional: si la lista está vacía...
        return 0.0
    return sum(numeros) / len(numeros)  # promedio = suma / cantidad


# LADRILLO función: decide el tipo de certificado según AGENTS.md.
def decidir_certificado(promedio, asistencia):
    # LADRILLO variables y tipos: promedio (float), asistencia (float) -> texto.
    # LADRILLO condicionales: los límites INCLUYEN el valor (>=).
    if asistencia < 80:
        return "Sin certificado"
    if promedio >= 70:
        return "Aprobación"
    return "Participación"


def leer_maestro(ruta):
    wb = load_workbook(ruta, data_only=True, read_only=True)
    ws = wb["Estudiantes"]
    maestro = {}  # LADRILLO variable tipo dict: clave (id, programa) -> datos.
    # LADRILLO bucle: recorre cada fila del Excel (salta el encabezado).
    for fila in list(ws.iter_rows(values_only=True))[1:]:
        ident, nombre, correo, programa, cohorte = fila
        maestro[(str(ident), programa)] = {
            "nombre": nombre, "correo": correo, "cohorte": cohorte,
        }
    return maestro


def leer_evaluaciones(ruta):
    wb = load_workbook(ruta, data_only=True, read_only=True)
    ws = wb["Evaluaciones"]
    notas = defaultdict(list)  # clave (id, programa) -> lista de notas
    asist = defaultdict(list)  # clave (id, programa) -> lista de asistencias
    modulos = set()
    todas_notas = []
    # LADRILLO bucle + condicionales para juntar datos por estudiante y programa.
    for fila in list(ws.iter_rows(values_only=True))[1:]:
        ident, programa, modulo, nota, asistencia, _cierre = fila
        clave = (str(ident), programa)
        notas[clave].append(float(nota))
        asist[clave].append(float(asistencia))
        modulos.add(modulo)
        todas_notas.append(float(nota))
    return notas, asist, modulos, todas_notas


def main():
    maestro = leer_maestro("INSUMOS/Maestro_Estudiantes.xlsx")
    notas, asist, modulos, todas_notas = leer_evaluaciones(
        "INSUMOS/Registro_Evaluaciones.xlsx")

    # Resumen por programa (maestro).
    por_programa = defaultdict(int)
    # LADRILLO bucle: cuenta estudiantes por programa.
    for (_ident, programa) in maestro:
        por_programa[programa] += 1
    print("== Resumen ==")
    print(f"Estudiantes en maestro: {len(maestro)}")
    for prog, cant in sorted(por_programa.items()):
        print(f"  {prog}: {cant}")
    print(f"Filas de evaluación: {sum(len(v) for v in notas.values())}")
    print(f"Grupos (id+programa) con notas: {len(notas)}")
    print(f"Módulos: {sorted(modulos)}")
    print(f"Rango de notas: {min(todas_notas)} - {max(todas_notas)}")

    # Cruce + cálculo + tabla ordenada por programa y nombre.
    print("\n== Tabla (programa, nombre, id, módulos, promedio, asistencia, tipo) ==")
    filas = []
    # LADRILLO bucle: por cada grupo con notas, calcula y clasifica.
    for clave in notas:
        promedio = calcular_promedio(notas[clave])  # LADRILLO función
        asistencia = calcular_promedio(asist[clave])
        tipo = decidir_certificado(promedio, asistencia)  # LADRILLO función
        nombre = maestro[clave]["nombre"] if clave in maestro else "(sin maestro)"
        filas.append((clave[1], nombre, clave[0],
                      len(notas[clave]), promedio, asistencia, tipo))
    filas.sort()  # ordena por programa y luego nombre
    for f in filas:
        print(f"{f[0]} | {f[1]} | {f[2]} | mod={f[3]} | "
              f"prom={f[4]:.2f} | asis={f[5]:.2f} | {f[6]}")

    # Inconsistencias.
    print("\n== Inconsistencias ==")
    sin_eval = [c for c in maestro if c not in notas]
    huerf = [c for c in notas if c not in maestro]
    print(f"Maestro SIN evaluación ({len(sin_eval)}): {sin_eval}")
    print(f"Evaluación SIN maestro ({len(huerf)}): {huerf}")


if __name__ == "__main__":
    main()
