# tests/test_app.py — pruebas automáticas con el cliente de Flask.
# Se ejecutan con: python -m pytest
# No abren el navegador ni el servidor: simulan un navegador contra la app importada.

import os
from io import BytesIO
from pathlib import Path
import re

from openpyxl import load_workbook
import pytest

# LADRILLO: variables de prueba ANTES de importar la app, para no depender del .env.
os.environ["ADMIN_CORREO"] = "admin@pruebas.cr"
os.environ["ADMIN_CLAVE"] = "clave-admin-pruebas"
os.environ["USUARIO_CORREO"] = "usuario@pruebas.cr"
os.environ["USUARIO_CLAVE"] = "clave-usuario-pruebas"
os.environ["SECRET_KEY"] = "clave-secreta-solo-para-pruebas"

import app as aplicacion  # noqa: E402  (va después de las variables de entorno)

# LADRILLO variables: rutas de los Excel de fábrica (solo se leen, nunca se tocan).
BASE = Path(__file__).resolve().parents[1]
RUTA_MAESTRO = BASE / "INSUMOS" / "Maestro_Estudiantes.xlsx"
RUTA_EVAL = BASE / "INSUMOS" / "Registro_Evaluaciones.xlsx"

ADMIN_CORREO = os.environ["ADMIN_CORREO"]
ADMIN_CLAVE = os.environ["ADMIN_CLAVE"]
USUARIO_CORREO = os.environ["USUARIO_CORREO"]
USUARIO_CLAVE = os.environ["USUARIO_CLAVE"]


# LADRILLO función: lee un Excel del disco y lo devuelve como "caja" en memoria.
def leer_bytes(ruta):
    with open(ruta, "rb") as archivo:
        return BytesIO(archivo.read())


# LADRILLO función: simula llenar el formulario de entrada.
def entrar(cliente, correo, clave):
    return cliente.post("/login", data={"correo": correo, "clave": clave})


# LADRILLO función: copia el Excel de evaluaciones en memoria y baja las notas.
def evaluaciones_con_notas_en_40():
    libro = load_workbook(RUTA_EVAL)  # LADRILLO: lee el original, no lo modifica
    hoja = libro["Evaluaciones"]
    # LADRILLO bucle: recorre las filas y busca a Camila Rojas Méndez.
    for fila in hoja.iter_rows(min_row=2):
        if str(fila[0].value) == "101230456":
            fila[3].value = 40  # columna Nota
    buffer = BytesIO()  # LADRILLO: guardamos la copia en memoria, no en disco
    libro.save(buffer)
    buffer.seek(0)
    return buffer


# LADRILLO fixture: cliente ya autenticado como administrador.
@pytest.fixture
def cliente_admin():
    cliente = aplicacion.app.test_client()
    respuesta = entrar(cliente, ADMIN_CORREO, ADMIN_CLAVE)
    assert respuesta.status_code == 302  # entró y lo mandaron al tablero
    return cliente


# LADRILLO función auxiliar: lee un contador del HTML (los números de arriba).
def contador(html, etiqueta):
    patron = rf"<strong>(\d+)</strong><span>{etiqueta}</span>"
    return int(re.search(patron, html).group(1))


def test_sin_sesion_redirige_a_login():
    """Simula a alguien que escribe la dirección de la app sin haber entrado."""
    cliente = aplicacion.app.test_client()
    respuesta = cliente.get("/")
    assert respuesta.status_code == 302
    assert "/login" in respuesta.headers["Location"]


def test_clave_incorrecta_no_entra():
    """Simula escribir bien el correo y mal la contraseña en la pantalla de entrada."""
    cliente = aplicacion.app.test_client()
    respuesta = cliente.post(
        "/login",
        data={"correo": ADMIN_CORREO, "clave": "clave-equivocada"},
    )
    assert respuesta.status_code == 200  # se queda en la pantalla de entrada
    assert "Correo o contraseña incorrectos" in respuesta.get_data(as_text=True)
    # Sigue afuera: la página principal lo vuelve a mandar al login.
    assert cliente.get("/").status_code == 302


def test_usuario_no_puede_cargar():
    """Simula a un usuario normal escribiendo /cargar a mano en la barra de direcciones."""
    cliente = aplicacion.app.test_client()
    entrar(cliente, USUARIO_CORREO, USUARIO_CLAVE)
    respuesta = cliente.get("/cargar")
    assert respuesta.status_code == 403
    assert "solo para administradores" in respuesta.get_data(as_text=True)


def test_excel_modificado_cambia_contadores(cliente_admin):
    """Simula subir los dos Excel con las notas de Camila en 40 desde el panel."""
    respuesta = cliente_admin.post(
        "/cargar",
        data={
            "accion": "pantalla",
            "maestro": (leer_bytes(RUTA_MAESTRO), "Maestro_Estudiantes.xlsx"),
            "evaluaciones": (
                evaluaciones_con_notas_en_40(),
                "Registro_Evaluaciones.xlsx",
            ),
        },
        content_type="multipart/form-data",
    )
    assert respuesta.status_code == 200
    html = respuesta.get_data(as_text=True)
    # LADRILLO: Camila pasa de aprobación a participación, sin tocar la asistencia.
    assert contador(html, "Aprobación") == 18
    assert contador(html, "Participación") == 4
    assert contador(html, "Sin certificado") == 1
    # LADRILLO: el aviso verde confirma que esos datos no se guardaron.
    assert "Viven solo en esta pantalla" in html
