# app.py — LA COCINA (backend con Flask).
# Toda la lógica va aquí. La página solo muestra lo que aquí se calcula.
# Para arrancar con doble clic usa iniciar_app.bat. Para detener: Ctrl+C.

from collections import defaultdict  # LADRILLO: caja para agrupar por clave
from datetime import date  # LADRILLO: fecha de emisión del certificado
from functools import wraps  # LADRILLO: conservar el nombre de la función decorada
import hmac  # LADRILLO: comparar contraseñas de forma segura
import os  # LADRILLO: leer variables de entorno (nunca claves en el código)
from pathlib import Path  # LADRILLO: rutas que funcionan desde cualquier carpeta
from dotenv import load_dotenv  # LADRILLO: leer el archivo .env en tu máquina
from flask import (  # LADRILLO: Flask = servidor web
    Flask, redirect, render_template, request, session, url_for,
)
from openpyxl import load_workbook  # LADRILLO: librería para leer Excel .xlsx

# LADRILLO variable: creamos la aplicación web.
app = Flask(__name__)

# LADRILLO variable: carpeta donde vive app.py (no depende de dónde se ejecuta).
BASE_DIR = Path(__file__).resolve().parent
RUTA_MAESTRO = BASE_DIR / "INSUMOS" / "Maestro_Estudiantes.xlsx"
RUTA_EVAL = BASE_DIR / "INSUMOS" / "Registro_Evaluaciones.xlsx"

# LADRILLO: carga el archivo .env de esta carpeta (si no existe, no pasa nada).
load_dotenv(BASE_DIR / ".env")

# LADRILLO variable: clave para firmar la cookie de sesión.
# En producción debe venir de una variable de entorno (SECRET_KEY).
app.secret_key = os.environ.get("SECRET_KEY", "clave-solo-para-demostracion")


# LADRILLO variable tipo dict: sigla de cada programa para armar el folio.
SIGLAS_PROGRAMA = {
    "Técnico en IA Aplicada": "IA",
    "Tecnico en IA Aplicada": "IA",  # variante sin tilde, por si cambia el Excel
    "Excel Avanzado para Negocios": "EXC",
}

# LADRILLO variable tipo dict: meses en español para escribir la fecha.
MESES = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio",
    7: "julio", 8: "agosto", 9: "setiembre", 10: "octubre", 11: "noviembre",
    12: "diciembre",
}

# LADRILLO variable tipo dict: título y texto según el tipo de certificado.
TITULOS_CERTIFICADO = {
    "Aprobación": "Certificado de Aprobación",
    "Participación": "Certificado de Participación",
}
TEXTOS_CERTIFICADO = {
    "Aprobación": "por haber completado y aprobado satisfactoriamente el programa",
    "Participación": "por su participación y cumplimiento de asistencia en el programa",
}

# LADRILLO variable tipo dict: usuarios de prueba si faltan las variables de entorno.
USUARIOS_DEMO = {
    "admin": {"correo": "admin@horizonte.cr", "clave": "admin1234"},
    "usuario": {"correo": "usuario@horizonte.cr", "clave": "usuario1234"},
}


# LADRILLO función: lee los dos usuarios desde el entorno (o usa los de prueba).
def cargar_usuarios():
    """Devuelve el diccionario de usuarios, si está en modo demo y los avisos."""
    usuarios = {}
    demo_visibles = []
    # LADRILLO bucle: un prefijo de variables por cada usuario permitido.
    for prefijo, rol in (("ADMIN", "admin"), ("USUARIO", "usuario")):
        correo = os.environ.get(f"{prefijo}_CORREO")
        clave = os.environ.get(f"{prefijo}_CLAVE")
        # LADRILLO condicional: si falta una variable, entramos en modo demostración.
        if not correo or not clave:
            demo = USUARIOS_DEMO[rol]
            correo, clave = demo["correo"], demo["clave"]
            demo_visibles.append({"correo": correo, "clave": clave, "rol": rol})
        # LADRILLO variable: el correo se guarda en minúsculas para comparar fácil.
        usuarios[correo.strip().lower()] = {"clave": clave, "rol": rol}
    return usuarios, bool(demo_visibles), demo_visibles


# LADRILLO variables: usuarios disponibles y aviso de modo demostración.
USUARIOS, MODO_DEMO, USUARIOS_DEMO_VISIBLES = cargar_usuarios()


# LADRILLO función decoradora: exige sesión iniciada antes de mostrar datos.
def requiere_login(vista):
    @wraps(vista)
    def envoltura(*args, **kwargs):
        # LADRILLO condicional: si no hay correo en la sesión, va a /login.
        if "correo" not in session:
            return redirect(url_for("login"))
        return vista(*args, **kwargs)
    return envoltura


# LADRILLO función decoradora: solo deja pasar al rol "admin".
def solo_admin(vista):
    @wraps(vista)
    def envoltura(*args, **kwargs):
        # LADRILLO condicional: cualquier rol distinto de admin recibe 403.
        if session.get("rol") != "admin":
            return render_template("sin_permiso.html"), 403
        return vista(*args, **kwargs)
    return envoltura



# LADRILLO función: calcula un promedio (suma / cantidad).
def calcular_promedio(numeros):
    # LADRILLO variables y tipos: numeros es lista de float, devuelve float.
    # LADRILLO condicional: si no hay datos, evita dividir entre cero.
    if not numeros:
        return 0.0
    return sum(numeros) / len(numeros)


# LADRILLO función: decide el certificado según AGENTS.md (límites inclusivos).
def decidir_certificado(promedio, asistencia):
    # LADRILLO variables y tipos: promedio (float), asistencia (float) -> texto.
    # LADRILLO condicionales: el orden importa (primero la asistencia).
    if asistencia < 80:
        return "Sin certificado"
    if promedio >= 70:
        return "Aprobación"
    return "Participación"


# LADRILLO función: escribe una fecha en español sencillo (ej. 21 de octubre de 2026).
def formatear_fecha(dia):
    # LADRILLO variable: día (date) -> texto.
    return f"{dia.day} de {MESES[dia.month]} de {dia.year}"


# LADRILLO función: arma el folio único de cada certificado.
def codigo_certificado(cohorte, programa, identificacion):
    # LADRILLO variables: cohorte (texto), programa (texto), id (texto) -> folio.
    sigla = SIGLAS_PROGRAMA.get(programa, "GEN")  # LADRILLO .get evita error
    cohorte_texto = cohorte if cohorte else "SIN-COHORTE"
    return f"AH-{cohorte_texto}-{sigla}-{identificacion}"


def cargar_maestro(ruta):
    """Lee el Excel oficial. Devuelve dict con clave (identificacion, programa)."""
    wb = load_workbook(ruta, data_only=True, read_only=True)
    ws = wb["Estudiantes"]
    maestro = {}
    # LADRILLO bucle: recorre cada fila del Excel (salta el encabezado).
    for fila in list(ws.iter_rows(values_only=True))[1:]:
        ident, nombre, correo, programa, cohorte = fila
        maestro[(str(ident), programa)] = {
            "nombre": nombre, "correo": correo, "cohorte": cohorte,
        }
    return maestro


def cargar_evaluaciones(ruta):
    """Lee notas. Devuelve notas y asistencias agrupadas por (id, programa)."""
    wb = load_workbook(ruta, data_only=True, read_only=True)
    ws = wb["Evaluaciones"]
    notas = defaultdict(list)
    asist = defaultdict(list)
    # LADRILLO bucle: junta las notas de cada estudiante y programa.
    for fila in list(ws.iter_rows(values_only=True))[1:]:
        ident, programa, _modulo, nota, asistencia, _cierre = fila
        clave = (str(ident), programa)
        notas[clave].append(float(nota))
        asist[clave].append(float(asistencia))
    return notas, asist


def calcular_resultados():
    """Cruza maestro + evaluaciones y prepara todo para la página."""
    maestro = cargar_maestro(RUTA_MAESTRO)
    notas, asist = cargar_evaluaciones(RUTA_EVAL)

    resultados = []
    # LADRILLO bucle: solo los del maestro generan certificado (el huérfano no).
    for clave, datos in maestro.items():
        if clave not in notas:
            continue  # LADRILLO condicional: sin notas -> va a inconsistencias
        promedio = calcular_promedio(notas[clave])  # LADRILLO función
        asistencia = calcular_promedio(asist[clave])
        tipo = decidir_certificado(promedio, asistencia)  # LADRILLO función
        resultados.append({
            "id": clave[0], "programa": clave[1], "nombre": datos["nombre"],
            "correo": datos["correo"], "cohorte": datos["cohorte"],
            "modulos": len(notas[clave]),
            "promedio": round(promedio, 2), "asistencia": round(asistencia, 2),
            "tipo": tipo,
            "folio": codigo_certificado(datos["cohorte"], clave[1], clave[0]),
        })
    # Orden: programa y luego nombre.
    resultados.sort(key=lambda r: (r["programa"], r["nombre"]))

    # LADRILLO condicionales para los contadores de arriba.
    n_aprob = sum(1 for r in resultados if r["tipo"] == "Aprobación")
    n_part = sum(1 for r in resultados if r["tipo"] == "Participación")
    n_sin = sum(1 for r in resultados if r["tipo"] == "Sin certificado")

    # Las 3 inconsistencias detectadas.
    sin_eval = [  # maestro SIN ninguna evaluación
        {"id": c[0], "programa": c[1], "nombre": d["nombre"]}
        for c, d in maestro.items() if c not in notas
    ]
    huerfanos = [  # evaluación SIN maestro (no genera certificado)
        {"id": c[0], "programa": c[1]}
        for c in notas if c not in maestro
    ]
    # LADRILLO bucle + condicional: detecta ceros (posible abandono).
    con_ceros = [
        {"id": c[0], "programa": c[1], "nombre": maestro[c]["nombre"]}
        for c in notas if c in maestro
        and (0 in notas[c] or 0 in asist[c])
    ]

    programas = sorted({r["programa"] for r in resultados})
    return resultados, n_aprob, n_part, n_sin, sin_eval, huerfanos, con_ceros, programas


# LADRILLO función: busca un solo estudiante por (identificación, programa).
def buscar_resultado(identificacion, programa):
    # LADRILLO variable: resultados es lista de dicts ya calculados.
    resultados, *_ = calcular_resultados()
    # LADRILLO bucle + condicional: compara como texto para no fallar con ceros.
    for r in resultados:
        if r["id"] == str(identificacion) and r["programa"] == programa:
            return r
    return None  # LADRILLO: si no existe, devolvemos None


# LADRILLO función con decorador: pantalla de entrada (GET muestra, POST valida).
@app.route("/login", methods=["GET", "POST"])
def login():
    # LADRILLO condicional: si ya hay sesión, no volvemos a pedir la contraseña.
    if session.get("correo"):
        return redirect(url_for("index"))

    error = None
    correo = ""
    if request.method == "POST":  # LADRILLO condicional: llegó el formulario
        correo = request.form.get("correo", "").strip().lower()
        clave = request.form.get("clave", "")
        datos = USUARIOS.get(correo)
        # LADRILLO condicional + hmac: compara la clave sin filtrar información.
        if datos and hmac.compare_digest(datos["clave"], clave):
            session.clear()  # LADRILLO: empezamos una sesión nueva y limpia
            session["correo"] = correo
            session["rol"] = datos["rol"]
            return redirect(url_for("index"))
        error = "Correo o contraseña incorrectos."

    return render_template(
        "login.html", error=error, correo=correo,
        modo_demo=MODO_DEMO, demo_visibles=USUARIOS_DEMO_VISIBLES,
    )


# LADRILLO función con decorador: cierra la sesión y vuelve a la entrada.
@app.route("/salir")
def salir():
    session.clear()  # LADRILLO: borra correo, rol y cualquier dato guardado
    return redirect(url_for("login"))


# LADRILLO función con decorador: atiende la página principal.
@app.route("/")
@requiere_login
def index():
    resultados, n_aprob, n_part, n_sin, sin_eval, huerfanos, con_ceros, programas = \
        calcular_resultados()
    # Filtro por programa (?programa=...). La página no decide, solo filtra vista.
    filtro = request.args.get("programa", "Todos")
    if filtro != "Todos":  # LADRILLO condicional
        resultados = [r for r in resultados if r["programa"] == filtro]
    return render_template(
        "index.html", resultados=resultados,
        n_aprob=n_aprob, n_part=n_part, n_sin=n_sin,
        n_certificados=n_aprob + n_part,
        sin_eval=sin_eval, huerfanos=huerfanos, con_ceros=con_ceros,
        programas=programas, filtro=filtro,
    )


# LADRILLO función: muestra el detalle académico de un estudiante.
@app.route("/estudiante/<identificacion>/<path:programa>")
@requiere_login
def estudiante(identificacion, programa):
    resultado = buscar_resultado(identificacion, programa)
    # LADRILLO condicional: si no existe, mensaje amable (no se cae la app).
    if resultado is None:
        return render_template(
            "mensaje.html",
            titulo="Estudiante no encontrado",
            mensaje=("No hay datos para la identificación "
                     f"{identificacion} en el programa {programa}."),
        ), 404
    return render_template("estudiante.html", r=resultado)


# LADRILLO función: emite el certificado solo si la regla lo permite.
@app.route("/certificado/<identificacion>/<path:programa>")
@requiere_login
def certificado(identificacion, programa):
    resultado = buscar_resultado(identificacion, programa)
    if resultado is None:
        return render_template(
            "mensaje.html",
            titulo="Estudiante no encontrado",
            mensaje=("No hay datos para la identificación "
                     f"{identificacion} en el programa {programa}."),
        ), 404
    # LADRILLO condicional: asistencia < 80 no genera certificado.
    if resultado["tipo"] == "Sin certificado":
        return render_template(
            "mensaje.html",
            titulo="Sin certificado",
            mensaje=(f"{resultado['nombre']} no cumple los requisitos "
                     f"(asistencia {resultado['asistencia']:.2f}%, "
                     "mínimo 80%)."),
        ), 403
    # LADRILLO: título y texto ya vienen decididos desde Python.
    return render_template(
        "certificado.html", r=resultado,
        titulo=resultado["tipo"],
        titulo_certificado=TITULOS_CERTIFICADO[resultado["tipo"]],
        texto_certificado=TEXTOS_CERTIFICADO[resultado["tipo"]],
        fecha_emision=formatear_fecha(date.today()),
    )


# LADRILLO función: junta TODOS los certificados en una sola página imprimible.
@app.route("/certificados")
@requiere_login
@solo_admin
def certificados():
    resultados, *_ = calcular_resultados()
    # LADRILLO bucle + condicional: solo quien tiene certificado entra al lote.
    aptos = [
        dict(r,
             titulo_certificado=TITULOS_CERTIFICADO[r["tipo"]],
             texto_certificado=TEXTOS_CERTIFICADO[r["tipo"]])
        for r in resultados if r["tipo"] != "Sin certificado"
    ]
    return render_template(
        "certificados.html", certificados=aptos,
        fecha_emision=formatear_fecha(date.today()),
    )


if __name__ == "__main__":
    # LADRILLO: arranca el servidor en http://localhost:5000
    app.run(debug=True)
